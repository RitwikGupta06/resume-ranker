import re
from pathlib import Path

import pymupdf
from docx import Document

SUPPORTED = {".pdf", ".docx", ".txt"}


def extract_text(path: str | Path) -> str:
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".pdf":
        raw = _pdf_text(path)
    elif ext == ".docx":
        raw = _docx_text(path)
    elif ext == ".txt":
        raw = path.read_text(encoding="utf-8", errors="ignore")
    else:
        raise ValueError(f"unsupported file type: {ext}")
    return clean(raw)


def _pdf_text(path: Path) -> str:
    # TODO: scanned PDFs come back empty here, add a tesseract fallback
    with pymupdf.open(path) as doc:
        return "\n".join(page.get_text() for page in doc)


def _docx_text(path: Path) -> str:
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    # a lot of resume templates put everything in tables
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def clean(text: str) -> str:
    text = text.replace("•", " ").replace("\xa0", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n\n", text)
    return text.strip()


def load_folder(folder: str | Path) -> dict[str, str]:
    """name -> text for every supported file in the folder."""
    out = {}
    for p in sorted(Path(folder).iterdir()):
        if p.suffix.lower() in SUPPORTED:
            out[p.stem] = extract_text(p)
    return out
