import pymupdf
import pytest
from docx import Document

from rater.parse import clean, extract_text, load_folder


def test_txt(tmp_path):
    f = tmp_path / "a.txt"
    f.write_text("Python   and\n\n\n\nSQL")
    assert extract_text(f) == "Python and\n\nSQL"


def test_pdf(tmp_path):
    f = tmp_path / "r.pdf"
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "Skills: Python, Docker")
    doc.save(f)
    assert "Python, Docker" in extract_text(f)


def test_docx_reads_tables(tmp_path):
    f = tmp_path / "r.docx"
    d = Document()
    d.add_paragraph("Jane Doe")
    d.add_table(rows=1, cols=1).cell(0, 0).text = "Kubernetes"
    d.save(str(f))
    text = extract_text(f)
    assert "Jane Doe" in text and "Kubernetes" in text


def test_unsupported(tmp_path):
    f = tmp_path / "r.odt"
    f.write_text("x")
    with pytest.raises(ValueError):
        extract_text(f)


def test_clean_bullets():
    assert clean("• Python\xa0dev") == "Python dev"


def test_load_folder_skips_other_files(tmp_path):
    (tmp_path / "a.txt").write_text("hi")
    (tmp_path / "notes.md").write_text("skip me")
    assert load_folder(tmp_path) == {"a": "hi"}
